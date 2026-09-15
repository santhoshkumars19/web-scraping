"""
app/services/export/excel_exporter.py

Excel (.xlsx) workbook generator using openpyxl.
Includes custom header styling, freeze panes, autofilter, column auto-sizing,
and text-safe cell formatting for phone numbers and codes.
"""

from __future__ import annotations

import io
from typing import Sequence
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TEXT_SAFE_HEADERS = {
    "Phone",
    "Alternate Phone",
    "WhatsApp",
    "Pincode",
}

# Theme palette matching LeadScout design
HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")  # Slate 900
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Calibri", size=10, color="0F172A")
THIN_BORDER_SIDE = Side(border_style="thin", color="E2E8F0")
DATA_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE,
)


def generate_excel_bytes(
    headers: list[str],
    rows: Sequence[dict[str, str]],
) -> bytes:
    """Generate a styled .xlsx workbook as raw bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    # 1. Write Header Row
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")

    # Set Header Row Height
    ws.row_dimensions[1].height = 26

    # 2. Write Data Rows
    for row_idx, row_data in enumerate(rows, start=2):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, header in enumerate(headers, start=1):
            val = row_data.get(header, "")
            cell = ws.cell(row=row_idx, column=col_idx)

            # Preserve string representation for phone numbers and pincodes
            if header in TEXT_SAFE_HEADERS:
                cell.data_type = "s"
                cell.number_format = "@"
                cell.value = str(val) if val is not None else ""
            else:
                cell.value = val

            cell.font = DATA_FONT
            cell.border = DATA_BORDER
            cell.alignment = Alignment(vertical="center")

    # 3. Freeze Top Header Row
    ws.freeze_panes = "A2"

    # 4. Enable Auto-filter across all columns
    if headers and rows:
        max_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{max_col_letter}{len(rows) + 1}"

    # 5. Compute readable column widths based on content length
    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(header))
        for row in rows:
            val_len = len(str(row.get(header, "")))
            if val_len > max_len:
                max_len = val_len

        # Clamp between 12 and 50 characters with extra safety padding
        adjusted_width = max(12, min(max_len + 3, 50))
        ws.column_dimensions[col_letter].width = adjusted_width

    # 6. Save to in-memory bytes
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
